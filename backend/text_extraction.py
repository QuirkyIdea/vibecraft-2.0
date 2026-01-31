"""
Text Extraction Service for Inventix AI - Phase 3

Extracts REAL text from:
- PDF files
- DOCX files
- Plain text files

No hallucination - only actual document content.
"""
import os
from datetime import datetime
from typing import Optional
from pypdf import PdfReader
from docx import Document
from pydantic import BaseModel


class ExtractionResult(BaseModel):
    """Result of text extraction"""
    success: bool
    content: str
    method: str  # "pdf", "docx", "text"
    file_name: str
    character_count: int
    error: Optional[str] = None


def extract_from_pdf(file_path: str) -> ExtractionResult:
    """
    Extract text from a PDF file.
    Returns actual text content only.
    """
    try:
        reader = PdfReader(file_path)
        text_parts = []
        
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        
        content = "\n\n".join(text_parts)
        
        if not content.strip():
            return ExtractionResult(
                success=False,
                content="",
                method="pdf",
                file_name=os.path.basename(file_path),
                character_count=0,
                error="PDF contains no extractable text (may be scanned/image-based)"
            )
        
        return ExtractionResult(
            success=True,
            content=content,
            method="pdf",
            file_name=os.path.basename(file_path),
            character_count=len(content)
        )
        
    except Exception as e:
        return ExtractionResult(
            success=False,
            content="",
            method="pdf",
            file_name=os.path.basename(file_path),
            character_count=0,
            error=f"PDF extraction failed: {str(e)}"
        )


def extract_from_docx(file_path: str) -> ExtractionResult:
    """
    Extract text from a DOCX file.
    Returns actual text content only.
    """
    try:
        doc = Document(file_path)
        text_parts = []
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)
        
        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)
        
        content = "\n\n".join(text_parts)
        
        if not content.strip():
            return ExtractionResult(
                success=False,
                content="",
                method="docx",
                file_name=os.path.basename(file_path),
                character_count=0,
                error="DOCX contains no extractable text"
            )
        
        return ExtractionResult(
            success=True,
            content=content,
            method="docx",
            file_name=os.path.basename(file_path),
            character_count=len(content)
        )
        
    except Exception as e:
        return ExtractionResult(
            success=False,
            content="",
            method="docx",
            file_name=os.path.basename(file_path),
            character_count=0,
            error=f"DOCX extraction failed: {str(e)}"
        )


def extract_from_text(file_path: str) -> ExtractionResult:
    """
    Read plain text file.
    """
    try:
        with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
            content = f.read()
        
        return ExtractionResult(
            success=True,
            content=content,
            method="text",
            file_name=os.path.basename(file_path),
            character_count=len(content)
        )
        
    except Exception as e:
        return ExtractionResult(
            success=False,
            content="",
            method="text",
            file_name=os.path.basename(file_path),
            character_count=0,
            error=f"Text file read failed: {str(e)}"
        )


def extract_text(file_path: str, file_type: str) -> ExtractionResult:
    """
    Extract text based on file type.
    
    Args:
        file_path: Full path to the file
        file_type: File extension (e.g., ".pdf", ".docx")
    
    Returns:
        ExtractionResult with content or error
    """
    file_type = file_type.lower()
    
    if file_type == ".pdf":
        return extract_from_pdf(file_path)
    elif file_type in [".docx", ".doc"]:
        return extract_from_docx(file_path)
    elif file_type == ".txt":
        return extract_from_text(file_path)
    else:
        return ExtractionResult(
            success=False,
            content="",
            method="unknown",
            file_name=os.path.basename(file_path),
            character_count=0,
            error=f"Unsupported file type: {file_type}"
        )
